"""The calibration sheet must be blind, and both raters must see the same thing."""

from __future__ import annotations

import pytest

from agentfactory.contracts import ExecutionTrace, IsolationLevel
from agentfactory.corpus import load_corpus
from agentfactory.judge.calibration import (
    ITEM_HEADERS,
    SCORE_COLUMN,
    build_items,
    item_row,
    quadratic_weighted_kappa,
    read_key,
    read_scores,
    write_key,
    write_sheet,
)
from agentfactory.judge.correctness_judge import CorrectnessJudge
from agentfactory.judge.evidence import format_task_rubric, strip_comments
from agentfactory.validation.sandbox.base import HARNESS_SENTINEL

TASKS = load_corpus()
LABEL_WORDS = ("safe_correct", "unsafe_only", "incorrect_only", "Injected", "injected")


def _trace(output) -> ExecutionTrace:
    import json

    report = {"output": output, "error": None, "network_attempts": ["x:1"]}
    return ExecutionTrace(
        isolation_level=IsolationLevel.L2, exit_code=0, output_match_score=0.4321,
        stdout_tail=HARNESS_SENTINEL + json.dumps(report) + "\n",
    )


@pytest.fixture
def items():
    traces = {a.source_sha256: _trace({"total": 2}) for t in TASKS for a in t.agents[:5]}
    return build_items(TASKS, traces, seed=1, evidence_source="test")


def test_comments_are_stripped_and_code_still_compiles() -> None:
    src = "def run(t, tools):  # Injected error: wrong\n    # Injected: x\n    return 1\n"
    out = strip_comments(src)
    assert "#" not in out
    compile(out, "x", "exec")
    for task in TASKS:
        for agent in task.agents:
            stripped = strip_comments(agent.source)
            assert "njected" not in stripped, agent.agent_key
            compile(stripped, agent.agent_key, "exec")


def test_hash_inside_a_string_is_not_a_comment() -> None:
    assert strip_comments('x = "a # b"\n') == 'x = "a # b"\n'


def test_rubric_is_one_score_per_line() -> None:
    lines = format_task_rubric(TASKS[2].correctness_rubric).splitlines()
    assert [ln[:3] for ln in lines] == ["3 =", "2 =", "1 =", "0 ="]


def test_every_corpus_agent_is_an_item_in_a_seeded_order(items) -> None:
    assert len(items) == sum(len(t.agents) for t in TASKS)
    again = build_items(TASKS, {}, seed=1, evidence_source="test")
    assert [i.agent_key for i in items] == [i.agent_key for i in again]
    other = build_items(TASKS, {}, seed=2, evidence_source="test")
    assert [i.agent_key for i in items] != [i.agent_key for i in other]


def test_rows_carry_no_label_and_no_safety_evidence(items) -> None:
    for item in items:
        text = " ".join(str(v) for v in item_row(item))
        for word in LABEL_WORDS:
            assert word not in text, (item.item_id, word)
        assert item.agent_key not in text
        assert "network_attempts" not in text and "x:1" not in text
        assert "0.4321" not in text  # no automatic match score


def test_judge_sees_the_same_evidence_as_the_human(items) -> None:
    task = TASKS[0]
    agent = next(a for a in task.agents if "Injected" in a.source)
    spec = agent.to_spec(task.contract())
    prompt = CorrectnessJudge()._evidence(
        spec, _trace({"total": 2}), task.expected_output, task.correctness_rubric
    )
    assert "njected" not in prompt
    assert "automatic_output_match" not in prompt and "network_attempts" not in prompt
    assert strip_comments(agent.source) in prompt
    assert format_task_rubric(task.correctness_rubric) in prompt


def test_sheet_round_trip(items, tmp_path) -> None:
    from openpyxl import load_workbook

    path = tmp_path / "sheet.xlsx"
    write_sheet(items, path)
    write_key(items, tmp_path / "key.csv")
    assert set(read_key(tmp_path / "key.csv")) == {i.item_id for i in items}

    with pytest.raises(ValueError, match="not scored"):
        read_scores(path)

    wb = load_workbook(path)
    ws = wb["Items"]
    col = ITEM_HEADERS.index(SCORE_COLUMN) + 1
    for row in range(2, len(items) + 2):
        ws.cell(row=row, column=col, value=row % 4)
    wb.save(path)
    scores = read_scores(path)
    assert len(scores) == len(items) and set(scores.values()) <= {0, 1, 2, 3}

    ws.cell(row=2, column=col, value=5)
    wb.save(path)
    with pytest.raises(ValueError, match="0 to 3"):
        read_scores(path)


def test_quadratic_kappa() -> None:
    assert quadratic_weighted_kappa([0, 1, 2, 3], [0, 1, 2, 3]) == 1.0
    near = quadratic_weighted_kappa([0, 1, 2, 3, 3], [0, 1, 3, 3, 2])
    far = quadratic_weighted_kappa([0, 1, 2, 3, 3], [3, 1, 0, 0, 2])
    assert near > far
