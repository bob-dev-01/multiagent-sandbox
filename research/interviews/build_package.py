"""Assemble the interview package as one Word document from the Markdown sources.

    python research/interviews/build_package.py

The Markdown files here are the record (committed before any interview, as the
plan requires); the .docx is for use in sessions and is written to
docs/interviews/, which is not in the repository. Handles the small Markdown
subset these files use: headings, paragraphs, bullet and numbered lists,
checkboxes, tables, **bold** and *italic*.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
OUT = HERE.parents[1] / "docs" / "interviews" / "Interview_Package_AgentFactory.docx"
PARTS = ["protocol.md", "consent.md", "guide_en.md", "guide_ru.md", "codebook.md",
         "joint_display.md"]
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def add_inline(paragraph, text: str) -> None:
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**"):
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`"):
            paragraph.add_run(piece[1:-1]).font.name = "Consolas"
        elif piece.startswith("*") and len(piece) > 1:
            paragraph.add_run(piece[1:-1]).italic = True
        else:
            paragraph.add_run(piece)


def add_table(doc, rows: list[str]) -> None:
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(set(c) <= set("-: ") for c in r)]
    table = doc.add_table(rows=len(cells), cols=len(cells[0]))
    table.style = "Table Grid"
    for i, row in enumerate(cells):
        for j, value in enumerate(row):
            para = table.cell(i, j).paragraphs[0]
            add_inline(para, value)
            if i == 0:
                for run in para.runs:
                    run.bold = True


def render(doc, text: str) -> None:
    # An indented line continues the list item above it.
    lines: list[str] = []
    for raw in text.splitlines():
        if raw.startswith("   ") and lines and re.match(r"^(- |\d+\. )", lines[-1]):
            lines[-1] += " " + raw.strip()
        else:
            lines.append(raw)
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            continue
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            doc.add_heading(line.lstrip("#").strip(), level=min(level, 3))
        elif line.strip() == "---":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif line.startswith("- [ ] "):
            add_inline(doc.add_paragraph(style="List Bullet"), "☐ " + line[6:])
        elif line.startswith("- "):
            add_inline(doc.add_paragraph(style="List Bullet"), line[2:])
        elif re.match(r"^\d+\. ", line):
            # Numbers kept from the text: Word's list numbering would run on
            # from one guide into the next.
            number, rest = line.split(". ", 1)
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Pt(18)
            para.paragraph_format.first_line_indent = Pt(-18)
            para.add_run(f"{number}.\t")
            add_inline(para, rest)
        elif line.strip():
            # Join wrapped lines of one paragraph.
            para_lines = [line]
            while (i + 1 < len(lines) and lines[i + 1].strip()
                   and not re.match(r"^(#|\||- |\d+\. |---)", lines[i + 1])):
                i += 1
                para_lines.append(lines[i])
            add_inline(doc.add_paragraph(), " ".join(s.strip() for s in para_lines))
        i += 1


def main() -> int:
    doc = Document()
    # python-docx's template omits the zoom percentage the schema requires.
    for zoom in doc.settings.element.iter(qn("w:zoom")):
        zoom.set(qn("w:percent"), "100")
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading("Agent Factory — practitioner interview package", level=0)
    for n, part in enumerate(PARTS):
        if n:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        render(doc, (HERE / part).read_text(encoding="utf-8"))
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
