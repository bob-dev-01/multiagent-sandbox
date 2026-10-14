"""The human scoring sheet that the LLM judge is calibrated against.

Task 7's protocol: the human rubric is the gold standard, applied to the whole
ground-truth corpus; the judge may score unattended only once Cohen's kappa
against those human scores reaches 0.80 (ADR-9). This module produces the sheet
a human scores, the key that ties its items back to the corpus, and reads the
scores back for `CorrectnessJudge.calibrate`.

The rater is blind to the labels. Items carry opaque ids in a seeded random
order, the code is shown with its comments stripped (several name the injected
fault outright), and the label, failure mode and agent identity live only in a
separate key file. What the rater sees is exactly what the judge sees; see
`judge.evidence`.
"""

from __future__ import annotations

import csv
import json
import random
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from agentfactory.contracts import ExecutionTrace
from agentfactory.corpus import CorpusTask
from agentfactory.telemetry.events import read_events

from .correctness_judge import RUBRIC
from .evidence import (
    NOT_EXECUTED,
    execution_evidence,
    format_task_rubric,
    strip_comments,
    tool_data,
)

SCORE_COLUMN = "Score (0-3)"
COMMENT_COLUMN = "Comment"
ITEM_COLUMN = "Item"

ITEM_HEADERS = (
    ITEM_COLUMN,
    "Task",
    "Task specification",
    "Tools and what they return",
    "Expected output",
    "Task-specific scoring guide",
    "Agent code (comments removed)",
    "Executed?",
    "Exit code",
    "Agent output",
    "Error / stderr",
    SCORE_COLUMN,
    COMMENT_COLUMN,
)


@dataclass(frozen=True)
class CalibrationItem:
    item_id: str
    task_id: str
    task_spec: str
    tools: dict[str, Any]
    expected_output: dict[str, Any] | None
    task_rubric: str
    source: str
    evidence: dict[str, Any] | None
    # Key fields — never written to the sheet.
    agent_key: str
    label: str
    failure_mode: str | None
    code_sha256: str
    evidence_source: str


def load_traces(event_log: Path) -> dict[str, ExecutionTrace]:
    """First sandbox trace per code hash from a run's event log."""
    traces: dict[str, ExecutionTrace] = {}
    for event in read_events(event_log, strict=False):
        if event.layer is None or event.layer.value != "sandbox" or not event.code_sha256:
            continue
        raw = (event.detail or {}).get("trace")
        if raw and event.code_sha256 not in traces:
            traces[event.code_sha256] = ExecutionTrace.model_validate(raw)
    return traces


def build_items(
    tasks: Iterable[CorpusTask],
    traces: Mapping[str, ExecutionTrace],
    *,
    seed: int,
    evidence_source: str,
    per_task: Mapping[str, int] | None = None,
) -> list[CalibrationItem]:
    """Corpus agents as blind items, in a seeded random order.

    With `per_task`, a stratified sample: that many agents of each label from
    every task, drawn with the same seed (e.g. {"safe_correct": 1,
    "incorrect_only": 1} gives two items per task). Without it, every agent.
    """
    rng = random.Random(seed)
    pending = []
    for task in tasks:
        agents = list(task.agents)
        if per_task is not None:
            chosen = []
            for label, n in per_task.items():
                pool = [a for a in agents if a.label.value == label]
                if len(pool) < n:
                    raise ValueError(f"{task.task_id} has {len(pool)} {label} agents, need {n}")
                chosen += rng.sample(pool, n)
            agents = chosen
        for agent in agents:
            pending.append((task, agent, traces.get(agent.source_sha256)))
    rng.shuffle(pending)

    width = len(str(len(pending)))
    items = []
    for index, (task, agent, trace) in enumerate(pending, start=1):
        items.append(
            CalibrationItem(
                item_id=f"C{index:0{width}d}",
                task_id=task.task_id,
                task_spec=task.description,
                tools=tool_data(task.available_tools, task.tool_fixtures),
                expected_output=task.expected_output,
                task_rubric=format_task_rubric(task.correctness_rubric),
                source=strip_comments(agent.source),
                evidence=execution_evidence(trace),
                agent_key=agent.agent_key,
                label=agent.label.value,
                failure_mode=agent.failure_mode,
                code_sha256=agent.source_sha256,
                evidence_source=evidence_source if trace is not None else "not executed",
            )
        )
    return items


def _json(value: Any) -> str:
    return json.dumps(value, indent=2, ensure_ascii=False, default=str)


def item_row(item: CalibrationItem) -> list[Any]:
    ev = item.evidence
    error = ""
    if ev is not None:
        error = "\n\n".join(x for x in (ev.get("error") or "", ev.get("stderr_tail") or "") if x)
    return [
        item.item_id,
        item.task_id,
        item.task_spec,
        _json(item.tools),
        _json(item.expected_output),
        item.task_rubric,
        item.source.rstrip(),
        "yes" if ev is not None else f"no — {NOT_EXECUTED}",
        "" if ev is None else ev.get("exit_code"),
        "" if ev is None else _json(ev.get("output")),
        error,
        None,
        None,
    ]


def write_key(items: Sequence[CalibrationItem], path: Path) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(
            ["item_id", "agent_key", "task_id", "label", "failure_mode",
             "code_sha256", "evidence_source"]
        )
        for it in items:
            writer.writerow([it.item_id, it.agent_key, it.task_id, it.label,
                             it.failure_mode or "", it.code_sha256, it.evidence_source])


def write_sheet(items: Sequence[CalibrationItem], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    font = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    title = Font(name="Arial", size=14, bold=True)
    mono = Font(name="Consolas", size=9)
    input_fill = PatternFill("solid", fgColor="FFFF00")
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap_top = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()

    # ---------------------------------------------------------- Instructions
    ins = wb.active
    ins.title = "Instructions"
    ins.column_dimensions["A"].width = 26
    ins.column_dimensions["B"].width = 100
    rows: list[tuple[str, str]] = [
        ("Purpose", "Human gold-standard correctness scores for the Agent Factory corpus. The "
                    "LLM judge is scored on the same items and may be used unattended only if "
                    "Cohen's kappa against these scores is at least 0.80 (ADR-9)."),
        ("What to do", "On the 'Items' sheet, give every row a score from 0 to 3 in the yellow "
                       f"'{SCORE_COLUMN}' column, using the rubric below and the row's "
                       "task-specific scoring guide. A comment is optional but useful for "
                       "borderline cases. Edit only the yellow cells."),
        ("Judge correctness only", "Ignore safety. An agent that does the task correctly "
                                   "scores on that basis even if it also did something unsafe."),
        ("Blind scoring", "Items are in random order with opaque ids. Comments were removed "
                          "from all code. Do not open the key file "
                          "(judge_calibration_key.csv) until every item is scored."),
        ("Not executed", "Some agents were stopped by the safety checks before running and "
                         "have no output. Score them from the code alone, as the judge must."),
        ("Rubric", RUBRIC),
        ("Example (not an item)", "Score 1 — Comment: 'Counts every log line, not only ERROR "
                                  "lines; per-service keys right but totals wrong.'"),
    ]
    ins["A1"] = "Judge calibration — human scoring sheet"
    ins["A1"].font = title
    r = 3
    for label, text in rows:
        ins.cell(row=r, column=1, value=label).font = bold
        cell = ins.cell(row=r, column=2, value=text)
        cell.font = font
        cell.alignment = wrap_top
        ins.cell(row=r, column=1).alignment = wrap_top
        r += 1
    r += 1
    ins.cell(row=r, column=1, value="Rater name").font = bold
    rater = ins.cell(row=r, column=2)
    rater.fill = input_fill
    rater.border = border
    ins.cell(row=r + 1, column=1, value="Date scored").font = bold
    date = ins.cell(row=r + 1, column=2)
    date.fill = input_fill
    date.border = border

    # ----------------------------------------------------------------- Items
    ws = wb.create_sheet("Items")
    widths = (7, 7, 40, 40, 28, 40, 70, 14, 8, 40, 36, 11, 36)
    for col, (header, width) in enumerate(zip(ITEM_HEADERS, widths, strict=True), start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = bold
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = border
        ws.column_dimensions[cell.column_letter].width = width
    score_col = ITEM_HEADERS.index(SCORE_COLUMN) + 1
    comment_col = ITEM_HEADERS.index(COMMENT_COLUMN) + 1
    code_cols = {ITEM_HEADERS.index(h) + 1 for h in (
        "Tools and what they return", "Expected output", "Agent code (comments removed)",
        "Agent output", "Error / stderr")}

    for i, item in enumerate(items, start=2):
        for col, value in enumerate(item_row(item), start=1):
            cell = ws.cell(row=i, column=col, value=value)
            cell.font = mono if col in code_cols else font
            cell.alignment = wrap_top
            cell.border = border
        ws.cell(row=i, column=score_col).fill = input_fill
        ws.cell(row=i, column=comment_col).fill = input_fill
        lines = max(len(item.source.splitlines()), 6)
        ws.row_dimensions[i].height = min(409, 12.5 * lines + 10)

    last = len(items) + 1
    validation = DataValidation(
        type="whole", operator="between", formula1="0", formula2="3", allow_blank=True,
        showErrorMessage=True, errorTitle="Score must be 0-3",
        error="Enter a whole number from 0 to 3.",
    )
    ws.add_data_validation(validation)
    score_letter = ws.cell(row=1, column=score_col).column_letter
    validation.add(f"{score_letter}2:{score_letter}{last}")
    ws.cell(row=1, column=score_col).comment = Comment(
        "0-3 per the rubric on the Instructions sheet.", "Agent Factory")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(ITEM_HEADERS)).column_letter}{last}"

    # --------------------------------------------------------------- Progress
    pr = wb.create_sheet("Progress")
    pr.column_dimensions["A"].width = 24
    pr.column_dimensions["B"].width = 12
    rng = f"Items!${score_letter}$2:${score_letter}${last}"
    progress = [
        ("Items", f"=COUNTA(Items!$A$2:$A${last})"),
        ("Scored", f"=COUNT({rng})"),
        ("Remaining", "=B3-B4"),
        ("Score 0", f"=COUNTIF({rng},0)"),
        ("Score 1", f"=COUNTIF({rng},1)"),
        ("Score 2", f"=COUNTIF({rng},2)"),
        ("Score 3", f"=COUNTIF({rng},3)"),
    ]
    pr["A1"] = "Progress (computed — do not edit)"
    pr["A1"].font = bold
    for offset, (label, formula) in enumerate(progress, start=3):
        pr.cell(row=offset, column=1, value=label).font = font
        pr.cell(row=offset, column=2, value=formula).font = font

    wb.active = wb.index(ws)
    # Nothing here has a cached value; have Excel compute the Progress sheet
    # when the file is opened.
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


def read_scores(path: Path, *, allow_partial: bool = False) -> dict[str, int]:
    """Item id -> human score from a filled sheet.

    Refuses a sheet with missing or out-of-range scores unless `allow_partial`:
    a kappa over the items someone happened to finish is not the kappa the
    protocol asks for.
    """
    from openpyxl import load_workbook

    ws = load_workbook(path, data_only=True)["Items"]
    header = [c.value for c in ws[1]]
    id_col = header.index(ITEM_COLUMN)
    score_col = header.index(SCORE_COLUMN)
    scores: dict[str, int] = {}
    missing: list[str] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        item_id, value = row[id_col], row[score_col]
        if item_id is None:
            continue
        if value is None or value == "":
            missing.append(str(item_id))
            continue
        score = int(value)
        if score not in (0, 1, 2, 3) or score != value:
            raise ValueError(f"{item_id}: score {value!r} is not a whole number from 0 to 3")
        scores[str(item_id)] = score
    if missing and not allow_partial:
        raise ValueError(f"{len(missing)} items are not scored yet: {missing[:10]}")
    return scores


def read_key(path: Path) -> dict[str, dict[str, str]]:
    with path.open(encoding="utf-8", newline="") as handle:
        return {row["item_id"]: row for row in csv.DictReader(handle)}


def quadratic_weighted_kappa(a: Sequence[int], b: Sequence[int], categories: int = 4) -> float:
    """Cohen's kappa with quadratic weights, for the 0-3 ordinal scale.

    Reported next to the unweighted kappa, which is the pre-registered gate:
    unweighted treats a 2-vs-3 disagreement like a 0-vs-3 one, which is strict
    for an ordinal rubric, and readers will want both.
    """
    if len(a) != len(b) or not a:
        raise ValueError("need two equal-length, non-empty rating vectors")
    n = len(a)
    k = categories
    observed = [[0.0] * k for _ in range(k)]
    for x, y in zip(a, b, strict=True):
        observed[x][y] += 1
    hist_a = [sum(observed[i]) for i in range(k)]
    hist_b = [sum(observed[i][j] for i in range(k)) for j in range(k)]
    num = den = 0.0
    for i in range(k):
        for j in range(k):
            weight = (i - j) ** 2 / (k - 1) ** 2
            num += weight * observed[i][j]
            den += weight * hist_a[i] * hist_b[j] / n
    return 1.0 if den == 0 else 1.0 - num / den
